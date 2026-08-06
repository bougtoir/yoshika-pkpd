"""Replace flattened Markdown table paragraphs with native Word tables."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


TABLES = [
    {
        "prefix": ": Table 1.",
        "caption": "Table 1. Proposed context-sensitive maximum dose framework.",
        "headers": ["Scenario", "Initial Compartment", "Expected C{sub:max}", "Dose Implication"],
        "rows": [
            ["Successful block", "BPT (V3)", "Low, delayed", "Higher dose may be safe"],
            ["Partial block", "Mixed (BPT + BRT)", "Intermediate", "Standard limit applies"],
            ["Failed block", "BRT (V2)", "Moderate-high, early", "Lower dose may be needed"],
            ["Intravascular injection", "Plasma (V1)", "Very high, immediate", "Traditional IV limits apply"],
            ["Epidural", "Depot (multi-pathway)", "Intermediate, delayed", "Depot model approximation"],
        ],
        "widths": [1.35, 1.55, 1.25, 2.15],
        "font_size": 9,
    },
    {
        "prefix": "Table 2.",
        "caption": "Table 2. Initial conditions for each selectable compartment.",
        "headers": [
            "Initial Compartment",
            "A{sub:depot}(0)",
            "A{sub:1}(0)",
            "A{sub:2}(0)",
            "A{sub:3}(0)",
        ],
        "rows": [
            ["Plasma (IV)", "0", "Dose", "0", "0"],
            ["BRT (failed block)", "0", "0", "Dose", "0"],
            ["BPT (successful block)", "0", "0", "0", "Dose"],
            ["Depot (fascial plane)", "Dose", "0", "0", "0"],
        ],
        "widths": [2.45, 0.95, 0.95, 0.95, 0.95],
        "font_size": 9,
    },
    {
        "prefix": ": Table 3.",
        "caption": (
            "Table 3. Built-in drug parameters with units and sources. Rate constants used "
            "internally are derived as k{sub:10}=CL/V{sub:1}, k{sub:12}=Q{sub:2}/V{sub:1}, "
            "k{sub:21}=Q{sub:2}/V{sub:2}, k{sub:13}=Q{sub:3}/V{sub:1}, and "
            "k{sub:31}=Q{sub:3}/V{sub:3}."
        ),
        "headers": ["Parameter (unit)", "Lidocaine", "Bupivacaine", "Ropivacaine", "Levobupivacaine"],
        "rows": [
            ["V{sub:1} (L)", "12.0", "8.9", "10.2", "9.5"],
            ["V{sub:2} (L)", "26.0", "22.6", "18.5", "24.0"],
            ["V{sub:3} (L)", "41.0", "51.3", "62.0", "48.0"],
            ["CL (L min{sup:−1})", "0.64", "0.47", "0.44", "0.50"],
            ["Q{sub:2} (L min{sup:−1})", "0.81", "0.60", "0.52", "0.55"],
            ["Q{sub:3} (L min{sup:−1})", "0.26", "0.17", "0.18", "0.16"],
            ["k{sub:e0} (min{sup:−1})", "0.12", "0.077", "0.065", "0.075"],
            ["k{sub:a} depot default (min{sup:−1})", "0.06", "0.04", "0.045", "0.042"],
            ["EC{sub:50} (mg L{sup:−1})", "3.0", "0.5", "0.8", "0.6"],
            ["γ (Hill, dimensionless)", "2.0", "3.0", "2.5", "2.8"],
            ["CNS toxic threshold (mg L{sup:−1})", "5.0", "2.0", "2.2", "2.5"],
            ["CV toxic threshold (mg L{sup:−1})", "10.0", "4.0", "5.3", "5.0"],
            ["Protein binding (fraction)", "0.65", "0.95", "0.94", "0.97"],
            ["pK{sub:a}", "7.9", "8.1", "8.1", "8.1"],
            ["Molecular weight (g mol{sup:−1})", "234.3", "288.4", "274.4", "288.4"],
        ],
        "widths": [2.25, 0.9, 1.05, 1.0, 1.25],
        "font_size": 8,
    },
    {
        "prefix": ": Table 4.",
        "caption": (
            "Table 4. External validation against published plasma concentration data. "
            "k{sub:a} was calibrated to the observed T{sub:max}; disposition parameters were "
            "unchanged. Fold error is predicted/observed C{sub:max}. The intravenous-input "
            "baseline predicts T{sub:max}=0 for every record."
        ),
        "headers": [
            "Study (route)",
            "Drug",
            "Dose (mg)",
            "Obs C{sub:max}",
            "Pred C{sub:max}",
            "Fold",
            "Obs T{sub:max} (min)",
            "IV C{sub:max}",
            "Ref",
        ],
        "rows": [
            ["Intercostal 0.75% (venous)", "ropivacaine", "150", "2.4", "4.69", "1.96", "11", "14.7", "[12]"],
            ["Intercostal 1.0% (venous)", "ropivacaine", "200", "2.5", "5.87", "2.35", "12", "19.6", "[12]"],
            ["Bilateral intercostal (venous)", "ropivacaine", "140", "1.1", "2.20", "2.00", "21", "13.7", "[13]"],
            ["Bilateral intercostal (venous)", "bupivacaine", "140", "0.9", "1.28", "1.42", "30", "15.7", "[13]"],
            ["TAP block (arterial)", "ropivacaine", "150", "1.83", "1.40", "0.76", "35", "14.7", "[14]"],
            ["Rectus sheath (arterial)", "ropivacaine", "150", "1.79", "1.03", "0.57", "53", "14.7", "[14]"],
            ["Brachial plexus (venous)", "ropivacaine", "190", "1.3", "1.30", "1.00", "53", "18.6", "[15]"],
            ["Axillary plexus (venous)", "ropivacaine", "175", "1.28", "1.21", "0.95", "52", "17.2", "[16]"],
            ["Axillary plexus (venous)", "bupivacaine", "175", "1.28", "1.06", "0.83", "58", "19.7", "[16]"],
            ["Epidural 2% (venous)", "lidocaine", "350", "2.3", "3.70", "1.61", "20", "29.2", "[17]"],
        ],
        "widths": [1.35, 0.72, 0.55, 0.62, 0.62, 0.43, 0.78, 0.58, 0.38],
        "font_size": 6.7,
    },
]

TOKEN_RE = re.compile(r"\{(sub|sup):([^}]+)\}")


def add_rich_text(paragraph, text: str, *, bold: bool = False, italic: bool = False,
                  size: float = 9, highlight: bool = False) -> None:
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.bold = bold
            run.italic = italic
            run.font.size = Pt(size)
            if highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        run = paragraph.add_run(match.group(2))
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if match.group(1) == "sub":
            run.font.subscript = True
        else:
            run.font.superscript = True
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_cell_text(cell, text: str, *, bold: bool, size: float, center: bool,
                  highlight: bool) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    add_rich_text(paragraph, text, bold=bold, size=size, highlight=highlight)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell, width_inches: float) -> None:
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def replace_flattened_tables(input_path: Path, output_path: Path, *, highlight: bool) -> None:
    document = Document(input_path)
    replaced = 0

    for spec in TABLES:
        paragraph = next((p for p in document.paragraphs if p.text.startswith(spec["prefix"])), None)
        if paragraph is None:
            raise RuntimeError(f"Flattened table paragraph not found: {spec['prefix']}")

        paragraph.clear()
        paragraph.style = document.styles["Caption"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(9)
        paragraph.paragraph_format.space_after = Pt(3)
        add_rich_text(
            paragraph,
            spec["caption"],
            italic=True,
            size=9,
            highlight=highlight,
        )

        table = document.add_table(rows=1, cols=len(spec["headers"]))
        table.style = document.styles["Table"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table._tbl.tblPr.append(OxmlElement("w:tblLayout"))
        table._tbl.tblPr[-1].set(qn("w:type"), "fixed")

        for index, text in enumerate(spec["headers"]):
            cell = table.rows[0].cells[index]
            set_cell_width(cell, spec["widths"][index])
            set_cell_text(
                cell,
                text,
                bold=True,
                size=spec["font_size"],
                center=index > 0,
                highlight=highlight,
            )
        set_repeat_header(table.rows[0])
        prevent_row_split(table.rows[0])

        for values in spec["rows"]:
            row = table.add_row()
            prevent_row_split(row)
            for index, text in enumerate(values):
                cell = row.cells[index]
                set_cell_width(cell, spec["widths"][index])
                set_cell_text(
                    cell,
                    text,
                    bold=False,
                    size=spec["font_size"],
                    center=index > 0,
                    highlight=highlight,
                )

        paragraph._p.addnext(table._tbl)
        replaced += 1

    if replaced != 4:
        raise RuntimeError(f"Expected 4 replacements, made {replaced}")
    document.save(output_path)
    print(f"Wrote {output_path} with {replaced} native tables")


def main() -> None:
    if len(sys.argv) not in (3, 4):
        raise SystemExit("Usage: python format_manuscript_tables.py INPUT.docx OUTPUT.docx [--highlight]")
    replace_flattened_tables(
        Path(sys.argv[1]),
        Path(sys.argv[2]),
        highlight=len(sys.argv) == 4 and sys.argv[3] == "--highlight",
    )


if __name__ == "__main__":
    main()
