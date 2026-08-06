"""Editable PowerPoint of the Software Impacts figures (one figure per slide).

Provided in addition to the inline figures in the manuscript docx so that each
figure can be edited or replaced independently.

Usage:
    python paper/software_impacts/build_figures_pptx.py
Output:
    paper/software_impacts/yoshika_software_impacts_figures.pptx
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
# Figures are regenerated into paper/figures by build_si.py / run_validation.py.
FIG = HERE.parent / "figures"
OUT = HERE / "yoshika_software_impacts_figures.pptx"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

FIGURES = [
    ("fig1_compartment_diagram_plasma.png",
     "Figure 1",
     "Three-compartment PK model with optional depot input. The initial "
     "compartment (plasma highlighted) is a user-selectable argument; "
     "disposition rate constants are shared across scenarios."),
    ("fig12_validation_cmax.png",
     "Figure 2",
     "External validation: observed versus predicted peak plasma concentration "
     "(depot model, ka calibrated to observed Tmax). Solid line, identity; "
     "shaded band, two-fold."),
]


def _add(prs: Presentation, img: Path, title: str, caption: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3),
                                  Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True

    cap = slide.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(12.3),
                                   Inches(0.8))
    cp = cap.text_frame.paragraphs[0]
    cp.text = caption
    cp.font.size = Pt(12)
    cap.text_frame.word_wrap = True

    with Image.open(img) as im:
        w, h = im.size
    avail_w, avail_h = Inches(12.0), Inches(5.3)
    ratio = min(avail_w / w, avail_h / h)
    dw, dh = int(w * ratio), int(h * ratio)
    left = int((SLIDE_W - dw) / 2)
    top = Inches(1.1)
    slide.shapes.add_picture(str(img), left, top, width=dw, height=dh)


def build() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for fname, title, caption in FIGURES:
        _add(prs, FIG / fname, title, caption)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(FIGURES)} slides)")


if __name__ == "__main__":
    build()
