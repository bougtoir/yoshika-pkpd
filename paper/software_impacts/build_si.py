"""Build the Software Impacts (Original Software Publication) manuscript for yoshika.

The manuscript follows the mandatory Software Impacts OSP template structure
(title / authors / abstract [~100 words] / keywords [<=6] / code-metadata table
C1-C9 / body <=3 pages / references), Software Impacts, Elsevier.

All quantitative results in the Impact section are computed at build time from
the bundled external-validation pipeline (paper/run_validation.py) rather than
hard-coded, so the article stays reproducible from the latest code.

Usage:
    python paper/software_impacts/build_si.py
Outputs:
    paper/software_impacts/yoshika_software_impacts.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

# allow `python paper/software_impacts/build_si.py` as well as `-m`
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import matplotlib
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

matplotlib.use("Agg")

import yoshika
from paper.run_validation import _fmt_stats, make_figures, run
from yoshika.plotting import plot_compartment_diagram

HERE = Path(__file__).resolve().parent
# Figures are regenerated from code + validation data into paper/figures so the
# article stays reproducible; nothing is committed as a stale static asset here.
FIG = HERE.parent / "figures"
OUT = HERE / "yoshika_software_impacts.docx"

FONT = "Arial"

# ── Author / contact metadata ────────────────────────────────────────────
AUTHOR = "Tatsuki Onishi"
ORCID = "0000-0001-7261-9062"
AFFIL = ("Data Science and AI Innovation Research Promotion Center, "
         "Shiga University, 1-1-1 Bamba, Hikone, Shiga 522-8522, Japan")
EMAIL = "bougtoir@gmail.com"
TEL = "+81-749-27-1023"

# Public code repository used for code-metadata C2 / C8.
REPO_URL = "https://github.com/bougtoir/yoshika-pkpd"
PYPI_URL = "https://pypi.org/project/yoshika/"

# ── References ─────────────────────────────────────────────────────────────
# All external references are generated in Vancouver style directly from the
# verified bibliography (paper/paper.bib) so DOIs/metadata cannot drift from the
# validated source. The enabling Array self-citation (published, with DOI) is
# supplied manually.
BIB_PATH = HERE.parent / "paper.bib"

MANUAL_REFS: dict[str, str] = {
    "onishi2026array": (
        "Onishi T. yoshika: a Python package for pharmacokinetic-pharmacodynamic "
        "simulation of local anesthetics with selectable initial compartment. "
        "Array 2026;31:101106. https://doi.org/10.1016/j.array.2026.101106."),
    "onishi2026zenodo": (
        "Onishi T. yoshika: a Python package for pharmacokinetic-pharmacodynamic "
        "simulation of local anesthetics with selectable initial compartment "
        "(v0.1.0). Zenodo; 2026. https://doi.org/10.5281/zenodo.21816214."),
}

_ACCENTS = {
    r'{\"o}': "o", r'{\"a}': "a", r'{\"u}': "u", r'{\"e}': "e",
    r"{\'e}": "e", r"{\'a}": "a", r"{\'o}": "o", r"{\'i}": "i",
    r"{\`e}": "e", r"{\^o}": "o",
}


def _clean(s: str) -> str:
    for k, v in _ACCENTS.items():
        s = s.replace(k, v)
    s = s.replace("{", "").replace("}", "")
    for esc in (r"\&", r"\%", r"\_", r"\#"):
        s = s.replace(esc, esc[1])
    return s.strip()


def _fmt_authors(raw: str) -> str:
    parts = [a.strip() for a in raw.split(" and ")]
    out = []
    trail = ""
    for a in parts:
        if a == "others":
            trail = ", et al"
            break
        if "," in a:
            last, given = [x.strip() for x in a.split(",", 1)]
        else:
            toks = a.split()
            last, given = toks[-1], " ".join(toks[:-1])
        initials = "".join(t[0] for t in re.split(r"[ .-]+", given) if t)
        out.append(f"{last} {initials}".strip())
    if len(out) > 6 and not trail:
        out = out[:6]
        trail = ", et al"
    return ", ".join(out) + trail


def parse_bib(path: Path) -> dict[str, dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    entries: dict[str, dict[str, str]] = {}
    for m in re.finditer(r"@\w+\{([^,]+),(.*?)\n\}", text, re.DOTALL):
        key = m.group(1).strip()
        fields: dict[str, str] = {}
        for fm in re.finditer(r"(\w+)\s*=\s*\{(.*?)\}\s*,?\s*(?:\n|$)",
                              m.group(2), re.DOTALL):
            fields[fm.group(1).lower()] = _clean(
                re.sub(r"\s+", " ", fm.group(2)))
        entries[key] = fields
    return entries


def format_ref(e: dict[str, str]) -> str:
    authors = _fmt_authors(e["author"])
    vol = e.get("volume", "")
    num = f"({e['number']})" if e.get("number") else ""
    pages = e.get("pages", "").replace("--", "-")
    loc = f"{vol}{num}:{pages}" if pages else f"{vol}{num}"
    s = f"{authors}. {e['title']}. {e['journal']} {e['year']};{loc}."
    if e.get("doi"):
        s += f" https://doi.org/{e['doi']}."
    return s


_BIB = parse_bib(BIB_PATH)


def _resolve_ref(key: str) -> str:
    if key in MANUAL_REFS:
        return MANUAL_REFS[key]
    if key in _BIB:
        return format_ref(_BIB[key])
    raise KeyError(f"reference not found in bib or manual list: {key}")


class Refs:
    """Assign [n] to citation keys in order of first appearance."""

    def __init__(self):
        self.order: list[str] = []

    def cite(self, *keys: str) -> str:
        nums = []
        for k in keys:
            _resolve_ref(k)  # validates key exists
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        return "[" + ",".join(str(n) for n in nums) + "]"

    def list(self) -> list[str]:
        return [_resolve_ref(k) for k in self.order]


# ── OMML (Word equation) helpers ────────────────────────────────────────────
# In-text formulae are emitted as native Word equations (OMML) rather than
# LaTeX or plain text, so they render and remain editable in Word's equation
# editor.
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def _m(tag: str):
    return OxmlElement(f"m:{tag}")


def _mrun(text: str):
    r = _m("r")
    t = _m("t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _msub(base: str, sub: str):
    ssub = _m("sSub")
    e = _m("e")
    e.append(_mrun(base))
    s = _m("sub")
    s.append(_mrun(sub))
    ssub.append(e)
    ssub.append(s)
    return ssub


def _mfrac(num, den):
    f = _m("f")
    n = _m("num")
    n.append(num)
    d = _m("den")
    d.append(den)
    f.append(n)
    f.append(d)
    return f


def math_concentration(par):
    """Append the inline Word equation C_i = A_i / V_i to a paragraph."""
    omath = _m("oMath")
    omath.append(_msub("C", "i"))
    omath.append(_mrun(" = "))
    omath.append(_mfrac(_msub("A", "i"), _msub("V", "i")))
    par._p.append(omath)


# ── docx helpers ──────────────────────────────────────────────────────────
SUB_RE = re.compile(r"~([^~]+)~")


def _set_font(run, size=10, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def rich(par, text, size=10, bold=False, italic=False):
    """Add text with ~x~ rendered as subscript."""
    pos = 0
    for m in SUB_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            _set_font(r, size, bold, italic)
        r = par.add_run(m.group(1))
        _set_font(r, size, bold, italic)
        r.font.subscript = True
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        _set_font(r, size, bold, italic)


def heading(doc, text, size=12):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_font(r, size, bold=True)
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p


def body(doc, text, size=10, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    rich(p, text, size=size, italic=italic)
    return p


def bullet(doc, lead, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    rich(p, lead, size=size, bold=True)
    rich(p, text, size=size)
    return p


def build():
    import matplotlib.pyplot as plt

    rows = run()

    # Regenerate the two figures used in the article from code + validation data.
    FIG.mkdir(exist_ok=True)
    make_figures(rows)  # writes fig12_validation_cmax.png (+ fig13) to paper/figures
    fig, _ = plot_compartment_diagram("plasma")
    fig.savefig(FIG / "fig1_compartment_diagram_plasma.png", dpi=300,
                bbox_inches="tight")
    plt.close(fig)

    stats = _fmt_stats(rows)
    gmfe = f"{stats['gmfe']:.2f}"
    within2 = f"{stats['within2_pct']:.0f}"
    pear = f"{stats['pearson_r']:.2f}"
    n_studies = len(rows)
    version = yoshika.__version__

    R = Refs()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    # ── Title ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("yoshika: A Python package for route-aware pharmacokinetic-"
                  "pharmacodynamic simulation of local anesthetics with a "
                  "selectable initial compartment")
    _set_font(r, 14, bold=True)

    # ── Authors ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, f"{AUTHOR} (ORCID {ORCID})", size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, AFFIL, size=9, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, f"Corresponding author: {AUTHOR}. Tel: {TEL}; e-mail: {EMAIL}",
         size=9)

    # ── Abstract (~100 words) ──
    heading(doc, "Abstract")
    body(doc,
         "yoshika is an open-source Python package for pharmacokinetic-"
         "pharmacodynamic (PKPD) simulation of local anesthetics in which the "
         "user chooses where the dose is first placed. Conventional simulators "
         "put the whole dose in plasma at time "
         "zero, which suits intravenous but not regional administration. yoshika "
         "instead starts the dose in plasma, vessel-rich tissue, vessel-poor "
         "tissue, or a first-order depot and compares concentration-time and "
         "effect-site profiles under shared disposition parameters. It ships "
         "parameters for four agents, a sigmoid-Emax effect-site model, plotting "
         "and tests, and a bundled external-validation pipeline against "
         f"{n_studies} published clinical datasets, so users can explore, "
         "reproducibly, how the administration route shifts peak exposure and "
         "its timing.")

    # ── Keywords ──
    heading(doc, "Keywords")
    body(doc, "Pharmacokinetics; Local anesthetics; Regional anesthesia; "
              "Compartment model; Simulation; Reproducibility")

    # ── Code metadata table ──
    heading(doc, "Code metadata")
    meta = [
        ("C1", "Current code version", f"v{version}"),
        ("C2", "Permanent link to code/repository used for this code version",
         REPO_URL),
        ("C3", "Permanent link to reproducible capsule",
         "Archived on Zenodo, DOI https://doi.org/10.5281/zenodo.21816213 "
         "(concept DOI, resolves to the latest version; v0.1.0 is "
         "https://doi.org/10.5281/zenodo.21816214). The archive bundles the "
         "package source, tests, external-validation pipeline and published "
         "reference values, and a self-contained capsule at "
         "paper/software_impacts/reproducible_capsule that reproduces the "
         "external-validation results and figures"),
        ("C4", "Legal code license", "MIT License"),
        ("C5", "Code versioning system used", "git"),
        ("C6", "Software code languages, tools and services used",
         "Python; NumPy, SciPy, Matplotlib, pandas"),
        ("C7", "Compilation requirements, operating environments and dependencies",
         "Python >= 3.9 (Linux/macOS/Windows); numpy>=1.21, scipy>=1.7, "
         "matplotlib>=3.5, pandas>=1.3. Install: pip install yoshika "
         f"({PYPI_URL})"),
        ("C8", "If available, link to developer documentation/manual",
         f"{REPO_URL}#readme"),
        ("C9", "Support email for questions", EMAIL),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for c, txt in zip(hdr, ["Nr", "Code metadata description", "Value"]):
        c.paragraphs[0].clear()
        rich(c.paragraphs[0], txt, size=9, bold=True)
    for nr, desc, val in meta:
        cells = table.add_row().cells
        for c, txt in zip(cells, [nr, desc, val]):
            c.paragraphs[0].clear()
            rich(c.paragraphs[0], txt, size=9)

    # ── Body ──
    heading(doc, "1. Motivation and significance")
    body(doc,
         "Local anesthetics used in regional anesthesia are deposited near "
         "nerves or in fascial planes, not into the bloodstream. Yet widely used "
         "pharmacokinetic (PK) simulation tools (for example STANPUMP, "
         "Tivatrainer, and Eleveld-model implementations) assume intravenous "
         "administration, placing the entire dose in the central plasma "
         f"compartment at time zero {R.cite('eleveld2018')}. General anesthesia "
         f"simulators such as the Python Anesthesia Simulator {R.cite('aubouinpairault2023')} "
         "solve compartmental PK systems but do not treat the site of initial "
         "drug deposition as a routine input. As a result the same model is "
         "applied to intravenous boluses and to tissue-deposited regional blocks, "
         "even though the route determines when and how high the plasma "
         "concentration peaks.")
    body(doc,
         "yoshika (Yielding Open Simulation of Hybrid Inter-disciplinary Kinetic "
         "Absorption) addresses this by making the initial compartment a "
         "user-selectable input. The dose can be placed in plasma "
         "(V~1~; intravenous), vessel-rich tissue (V~2~; e.g. a failed or "
         "intravascular block), vessel-poor tissue (V~3~; e.g. a successful "
         "nerve block), or a first-order depot (e.g. a fascial-plane block or "
         "infiltration). All scenarios share the same disposition parameters, so "
         "differences arise solely from the modeled route. The equations "
         "themselves are a conventional three-compartment mammillary model with "
         "an optional depot; yoshika's contribution is not a new equation but a "
         "reproducible, openly available way to make the drug-deposition "
         "assumption explicit and to compare routes side by side.")

    heading(doc, "2. Software description")
    body(doc,
         "yoshika is a small, dependency-light Python package (NumPy, SciPy, "
         f"Matplotlib, pandas) {R.cite('virtanen2020')} organized into focused "
         "modules:", space_after=4)
    bullet(doc, "compartments: ",
           "definitions of PLASMA, BRT (vessel-rich tissue), BPT (vessel-poor "
           "tissue) and DEPOT and the initial-condition logic selected by the "
           "user.")
    bullet(doc, "drugs: ",
           "a parameter database for lidocaine, bupivacaine, ropivacaine and "
           "levobupivacaine, drawn from the clinical PK literature "
           f"{R.cite('tucker1979', 'burm1989')}, plus an API to add custom drugs.")
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    rich(p, "model: ", bold=True)
    rich(p, "the three-compartment PK system with an optional depot input, solved "
            "with SciPy solve_ivp (RK45). A concentration is the amount in a "
            "compartment divided by its volume, ")
    math_concentration(p)
    rich(p, ".")
    bullet(doc, "pd: ",
           "an effect-site compartment (k~e0~) with a sigmoid-Emax response.")
    bullet(doc, "simulator / plotting / utils: ",
           "a high-level API for single runs and multi-scenario comparison, "
           "Matplotlib visualizations, and helpers (AUC, half-life, units).")
    body(doc,
         "The initial compartment sets which state receives the dose at t=0 "
         "(depot introduces first-order absorption at rate k~a~); the shared "
         "rate constants k~10~, k~12~, k~21~, k~13~, k~31~ are unchanged across "
         "scenarios. A unit-test suite covers the model, drug database, PD model "
         "and utilities. The package is released under the MIT license, "
         f"versioned with git, installable from PyPI ({PYPI_URL}), and archived "
         f"on Zenodo {R.cite('onishi2026zenodo')}.")
    body(doc,
         f"Fig. {_fig(1)} shows the model with the dose placed in plasma; the "
         "same diagram with any other compartment highlighted is produced by "
         "changing a single argument.")
    _add_figure(doc, FIG / "fig1_compartment_diagram_plasma.png",
                f"Fig. {_fig(1)}. Three-compartment PK model with optional depot "
                "input. The initial compartment (plasma highlighted) is a "
                "user-selectable argument; disposition rate constants are shared "
                "across scenarios.")

    heading(doc, "3. Illustrative example")
    body(doc,
         "Comparing routes for bupivacaine 150 mg in a 70 kg patient takes a few "
         "lines; the same disposition parameters are reused across scenarios:",
         space_after=4)
    _code(doc,
          "from yoshika import Simulator, Drug\n"
          "sim = Simulator(drug=Drug.BUPIVACAINE, dose_mg=150, weight_kg=70)\n"
          "results = sim.compare()            # plasma / BRT / BPT scenarios\n"
          "print(Simulator.summary_table(results))")

    heading(doc, "4. Impact")
    body(doc,
         "yoshika enables research questions that intravenous-input simulators "
         "make awkward to ask: how much does the assumed deposition route change "
         "predicted peak concentration (C~max~) and time-to-peak (T~max~), and "
         "how well does a standard model reproduce concentrations actually "
         "observed after regional administration? To answer the second question "
         "the package bundles a reproducible external-validation pipeline. Using "
         "the built-in disposition parameters unchanged and calibrating only the "
         "depot absorption rate k~a~ to each study's reported T~max~ (one degree "
         f"of freedom per study), yoshika was compared against {n_studies} "
         "published clinical PK records spanning intercostal "
         f"{R.cite('behnke2002', 'kopacz1994')}, fascial-plane "
         f"{R.cite('murouchi2015')}, peripheral-nerve "
         f"{R.cite('hickey1990', 'vainionpaa1995')}, and epidural "
         f"{R.cite('inoue1985')} routes.")
    body(doc,
         f"Across the {n_studies} records the geometric mean fold error of "
         f"predicted versus observed C~max~ was {gmfe}, with {within2}% of "
         f"predictions within two-fold of observation and a Pearson correlation "
         f"of r = {pear} (Fig. {_fig(2)}). By contrast, the intravenous-input "
         "baseline overestimated C~max~ by roughly an order of magnitude and "
         "predicted an instantaneous peak (T~max~ = 0) for every study, whereas "
         "observed peaks occurred tens of minutes after injection. The tool thus "
         "reproduces the correct order of magnitude and the qualitative timing of "
         "systemic exposure across routes, while not being accurate enough for "
         "quantitative dosing decisions.")
    _add_figure(doc, FIG / "fig12_validation_cmax.png",
                f"Fig. {_fig(2)}. External validation: observed versus predicted "
                "peak plasma concentration (depot model, k~a~ calibrated to "
                "observed T~max~). Solid line, identity; shaded band, two-fold. "
                f"{within2}% of predictions lie within two-fold (geometric mean "
                f"fold error {gmfe}; r = {pear}).")
    body(doc,
         "These results supported a peer-reviewed study of route-dependent "
         f"systemic-exposure and monitoring-window hypotheses {R.cite('onishi2026array')}, "
         "which is the scholarly publication enabled by the software. Beyond that "
         "study, the package is intended as an exploratory and educational tool: "
         "it lets clinicians and trainees visualize how injection site and "
         "absorption assumptions shift peak exposure and its timing, supports "
         "reconsidering how intravenous-derived mg/kg limits are applied to "
         f"regional anesthesia {R.cite('neal2018')}, and provides a small, "
         "route-aware PKPD component that could be embedded in anesthesia "
         "information management systems. The bundled dataset and pipeline give a "
         "reproducible baseline for further comparison.")

    heading(doc, "5. Limitations and future work")
    body(doc,
         "V~2~ and V~3~ are fitted lumped compartments, not anatomical spaces, "
         "and the disposition parameters are intravenous-derived; transferring "
         "them to peripheral injection sites is an approximation supported for "
         "order-of-magnitude and timing fidelity but not for quantitative "
         "dosing. Spinal (intrathecal) administration is not modeled. The "
         "validation relies on summary C~max~/T~max~ from the literature rather "
         "than full digitized curves. Planned work includes additional agents "
         "and population-PK support, richer depot/absorption models, and "
         "prospective linkage of an objective block-success signal to measured "
         "plasma concentrations.")

    heading(doc, "Acknowledgements")
    body(doc, "None.")

    heading(doc, "Declaration of competing interest")
    body(doc,
         "The author declares no known competing financial interests or personal "
         "relationships that could have appeared to influence the work reported "
         "in this paper.")

    heading(doc, "References")
    for i, ref in enumerate(R.list(), start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        rich(p, f"[{i}] {ref}", size=9)

    doc.save(OUT)

    # Highlights (3-5 bullets, each <= 85 chars) with computed stats.
    highlights = [
        "Open-source Python PKPD package for local anesthetics with selectable route",
        "The user selects the initial compartment: plasma, tissue, or depot",
        "Shared disposition parameters isolate the effect of the administration route",
        f"Validated vs {n_studies} published datasets: {within2}% of Cmax within "
        f"two-fold (GMFE {gmfe})",
        "IV-input assumption overstates Cmax ~10-fold and mis-times the peak (Tmax=0)",
    ]
    hl_path = HERE / "highlights.txt"
    hl_path.write_text("Highlights\n\n" + "\n".join(f"- {h}" for h in highlights)
                       + "\n", encoding="utf-8")
    over = [h for h in highlights if len(h) > 85]
    if over:
        raise ValueError(f"highlight exceeds 85 chars: {over}")

    # Report
    n_words = len(re.sub(r"[^\w\s]", " ",
                         " ".join(p.text for p in doc.paragraphs)).split())
    print(f"Wrote {OUT}")
    print(f"  version v{version}; {len(R.list())} references; "
          f"validation n={n_studies}, GMFE={gmfe}, within2={within2}%, r={pear}")
    print(f"  approx. total words in doc: {n_words}")


# figure numbering by first appearance
_FIG_ORDER: list[int] = []


def _fig(n: int) -> int:
    if n not in _FIG_ORDER:
        _FIG_ORDER.append(n)
    return _FIG_ORDER.index(n) + 1


def _add_figure(doc, path: Path, caption: str):
    from docx.shared import Inches
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(str(path), width=Inches(4.6))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(8)
    rich(cap, caption, size=9, italic=True)


def _code(doc, code: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(code)
    r.font.name = "Consolas"
    r.font.size = Pt(9)


if __name__ == "__main__":
    build()
