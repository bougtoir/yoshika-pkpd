"""Build the revision (R1) package for SIMPAC-D-26-00173.

Outputs (paper/software_impacts/):
  response_to_reviewers.docx              point-by-point response letter (from .md)
  yoshika_software_impacts_highlighted.docx
        revised manuscript with paragraphs that differ from the submitted
        version (yoshika_software_impacts_submitted_r0.docx) highlighted yellow

Usage:
    python paper/software_impacts/build_si.py
    python paper/software_impacts/build_revision_package.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
FONT = "Arial"

RESPONSE_MD = HERE / "response_to_reviewers.md"
RESPONSE_DOCX = HERE / "response_to_reviewers.docx"
SUBMITTED = HERE / "yoshika_software_impacts_submitted_r0.docx"
REVISED = HERE / "yoshika_software_impacts.docx"
HIGHLIGHTED = HERE / "yoshika_software_impacts_highlighted.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def _runs(par, text: str, size: int = 10):
    for tok in INLINE.split(text):
        if not tok:
            continue
        bold = italic = mono = False
        if tok.startswith("**") and tok.endswith("**"):
            tok, bold = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok, mono = tok[1:-1], True
        elif tok.startswith("*") and tok.endswith("*"):
            tok, italic = tok[1:-1], True
        r = par.add_run(tok)
        r.font.name = "Consolas" if mono else FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic


def build_response():
    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10)
    lines = RESPONSE_MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            _runs(p, f"**{line[2:]}**", size=14)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            _runs(p, f"**{line[3:]}**", size=12)
        elif line.startswith("> "):
            quote = []
            while i < len(lines) and lines[i].startswith("> "):
                quote.append(lines[i][2:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.space_after = Pt(6)
            _runs(p, " ".join(quote), size=9)
            for r in p.runs:
                r.font.italic = True
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            _runs(p, line)
        i += 1
    doc.save(RESPONSE_DOCX)
    print(f"Wrote {RESPONSE_DOCX}")


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def build_highlighted():
    old = {_norm(p.text) for p in Document(SUBMITTED).paragraphs if p.text.strip()}
    for t in Document(SUBMITTED).tables:
        for row in t.rows:
            for c in row.cells:
                old.add(_norm(c.text))
    doc = Document(REVISED)
    n = 0
    for p in doc.paragraphs:
        if p.text.strip() and _norm(p.text) not in old:
            for r in p.runs:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            n += 1
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if c.text.strip() and _norm(c.text) not in old:
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    n += 1
    doc.save(HIGHLIGHTED)
    print(f"Wrote {HIGHLIGHTED} ({n} changed paragraphs/cells highlighted)")


if __name__ == "__main__":
    build_response()
    build_highlighted()
