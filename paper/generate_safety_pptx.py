#!/usr/bin/env python3
"""Generate editable PPTX of the manuscript figures (one figure per slide).

Captions and figure numbers match the revised Array manuscript. Images are
scaled to preserve aspect ratio. All text is in English and editable.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

FIG_DIR = Path(__file__).resolve().parent / "figures"
PPTX_PATH = FIG_DIR / "manuscript_figures.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
AREA_LEFT = Inches(0.5)
AREA_TOP = Inches(1.0)
AREA_W = Inches(12.3)
AREA_H = Inches(5.0)

FIGURES = [
    ("fig0_graphical_abstract.png", "Graphical Abstract",
     "Selectable initial compartment -> order-of-magnitude effect on predicted peak and "
     "timing -> external validation against published data (90% within two-fold; GMFE 1.51; "
     "r = 0.81)."),
    ("fig1_compartment_diagram_plasma.png", "Figure 1: Three-compartment model",
     "Three-compartment model with selectable initial compartment (plasma highlighted)."),
    ("fig2_bupivacaine_comparison.png", "Figure 2: Bupivacaine by initial compartment",
     "Bupivacaine 150 mg: plasma concentration by initial compartment with toxicity thresholds."),
    ("fig3_all_drugs_comparison.png", "Figure 3: Four local anesthetics",
     "Plasma concentration comparison for all four local anesthetics by initial compartment."),
    ("fig4_bupivacaine_effect.png", "Figure 4: Effect-site response",
     "Bupivacaine: effect-site response by initial compartment."),
    ("fig5_bupivacaine_bpt_full.png", "Figure 5: Full compartment profile",
     "Bupivacaine 150 mg from vessel-poor tissue (BPT): full compartment concentration profile."),
    ("fig6_summary_table.png", "Figure 6: PKPD summary",
     "Bupivacaine 150 mg: PKPD summary by initial compartment."),
    ("fig7_safety_panel.png", "Figure 7: Route-aware scenarios",
     "Plasma concentration-time curves: intravenous-input model (red) vs depot/fascial plane "
     "(blue) and vessel-poor tissue/BPT (green dashed). Simulation scenarios; shared rate "
     "constants, only the initial compartment differs."),
    ("fig8_tmax_comparison.png", "Figure 8: Time to peak (Tmax)",
     "The intravenous-input model predicts Tmax = 0 for all drugs. Route-aware scenarios "
     "place the simulated peak tens of minutes later. Simulation outputs, not measured values."),
    ("fig9_monitoring_timeline.png", "Figure 9: Monitoring window",
     "Simulated peak timing relative to discharge. Diamond markers = Tmax under each scenario. "
     "Simulation scenarios, not clinically proven predictions."),
    ("fig10_cmax_toxicity.png", "Figure 10: Cmax vs toxicity thresholds",
     "The intravenous-input model predicts Cmax well above thresholds for non-intravenous "
     "routes; depot and vessel-poor-tissue scenarios predict lower values closer to published "
     "studies. Simulation outputs."),
    ("fig11_safety_summary_table.png", "Figure 11: Scenario summary",
     "Rate constants shared across scenarios; only the initial compartment differs. "
     "Pink = intravenous-input; blue = depot/fascial plane; green = vessel-poor tissue."),
    ("fig12_validation_cmax.png", "Figure 12: Validation — observed vs predicted Cmax",
     "Observed vs predicted Cmax (depot model, ka calibrated to observed Tmax). Solid line = "
     "identity; shaded band = two-fold. 90% within two-fold (GMFE 1.51; r = 0.81). Circle = "
     "venous, triangle = arterial sampling."),
    ("fig13_validation_tmax.png", "Figure 13: Validation — observed vs predicted Tmax",
     "The depot model reproduces observed 11-58 min peaks; the intravenous-input baseline "
     "predicts Tmax = 0 for every study (red crosses on the horizontal axis)."),
    ("fig14_validation_curves.png", "Figure 14: Validation — concentration-time curves",
     "Representative predicted curves (yoshika depot, blue; intravenous-input baseline, red "
     "dashed) against observed peak Cmax +/- SD (black squares) for intercostal, TAP, axillary "
     "plexus, and epidural cases."),
]


def _scaled(img_path: Path):
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / h
    area_ar = AREA_W / AREA_H
    if ar >= area_ar:
        width = AREA_W
        height = int(AREA_W / ar)
        left = AREA_LEFT
        top = int(AREA_TOP + (AREA_H - height) / 2)
    else:
        height = AREA_H
        width = int(AREA_H * ar)
        top = AREA_TOP
        left = int(AREA_LEFT + (AREA_W - width) / 2)
    return int(left), int(top), int(width), int(height)


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for fname, title, caption in FIGURES:
        img = FIG_DIR / fname
        if not img.exists():
            print(f"  WARNING: {img} not found, skipping")
            continue
        slide = prs.slides.add_slide(blank)

        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        left, top, width, height = _scaled(img)
        slide.shapes.add_picture(str(img), left, top, width, height)

        cb = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(12.3), Inches(1.1))
        cf = cb.text_frame
        cf.word_wrap = True
        cp = cf.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(12)
        cp.font.italic = True
        cp.alignment = PP_ALIGN.LEFT

    prs.save(str(PPTX_PATH))
    print(f"PPTX saved: {PPTX_PATH} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
